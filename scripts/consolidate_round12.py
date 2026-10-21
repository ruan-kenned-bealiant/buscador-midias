"""Consolida a rodada 12 offline, sem buscar páginas nem alterar o estado da coleta."""

from __future__ import annotations

import csv
import hashlib
import json
import sqlite3
from collections import Counter, defaultdict
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from buscador_midias.domain.url import canonicalize_url, url_hash
from buscador_midias.sites.loader import load_sites

ROOT = Path(__file__).resolve().parents[1]
ROUND = ROOT / 'var/ciclo-busca-20261019T005450Z'
OUTPUT = ROOT / 'var/analise-rodada-12/etapa-1'
LABELS = {
    'acessivel_com_candidatas': 'XML acessível, com links candidatos na janela',
    'acessivel_sem_candidatas': 'XML acessível, sem links candidatos datados na janela',
    'falha': 'Sem XML válido comprovado nesta execução',
    'parcial': 'Varredura parcial, encerrada por solicitação do usuário',
    'sem_detalhamento': 'Sem evidência individual suficiente para concluir',
    'descoberta': 'Página/robots de descoberta; não é RSS/sitemap',
}
CRITERIA = [
    'Etapa 1: consolidação local; não acessa sites, não usa IA e não seleciona mecanismos.',
    'Links candidatos não são notícias textuais validadas. '
    'Não houve leitura de artigos nesta etapa.',
    'A janela é fixa. Data de modificação não comprova publicação recente.',
    'Acessível significa RSS/sitemap interpretado, não acesso confirmado à página da notícia.',
    'Ausência de candidatas na janela não é erro nem justificativa de exclusão.',
    'As 163 fontes anteriores à fila preservam resultados agregados; '
    'não se inferem testes de raízes.',
    'Filho encontrado como origem direta não comprova leitura de todos os mecanismos ancestrais.',
    'URLs do CSV são deduplicadas globalmente pelo coletor; '
    'a atribuição é à primeira fonte/origem.',
    'Contagem por raiz inclui descendentes e pode se sobrepor; não somar como notícias únicas.',
    'Tarefas contam processamento por raiz, não chamadas HTTP: o cache compartilha respostas.',
    'Fila encerrada não garante cobertura integral: '
    'datas antigas, escopo e profundidade filtram filhos.',
    'O NZ Herald permanece parcial; os documentos pendentes foram preservados e não executados.',
    'O inventário original e todos os arquivos de entrada são preservados.',
]


def read_json(path: Path) -> Any:
    """Lê um artefato local."""
    return json.loads(path.read_text(encoding='utf-8'))


def write_csv(name: str, rows: list[dict[str, Any]]) -> None:
    """Exporta CSV com BOM para abertura em planilhas."""
    with (OUTPUT / name).open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def source_category(partial: bool, status: str, candidates: int) -> str:
    """Mantém incompletude independente da existência de resultados."""
    if partial:
        return 'parcial'
    if status == 'ok':
        return 'acessivel_com_candidatas' if candidates else 'acessivel_sem_candidatas'
    return 'falha' if status else 'sem_detalhamento'


def workbook(tables: dict[str, list[dict[str, Any]]]) -> None:
    """Gera planilha sem interpretar textos externos como fórmulas."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.utils import get_column_letter

    book = Workbook(write_only=True)
    for name, rows in tables.items():
        sheet = book.create_sheet(name)
        sheet.freeze_panes = 'A2'
        header = list(rows[0])
        for index, column in enumerate(header, 1):
            width = 65 if column in ('url', 'url_origem', 'exemplos', 'exemplo') else 22
            if column in ('detalhe', 'diagnostico', 'criterio', 'evidencia_documental'):
                width = 65
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.append(header)
        for row in rows:
            values = []
            for key in header:
                value = row[key]
                if isinstance(value, str):
                    cell = WriteOnlyCell(sheet, value=ILLEGAL_CHARACTERS_RE.sub('', value[:32767]))
                    cell.data_type = 's'
                    values.append(cell)
                else:
                    values.append(value)
            sheet.append(values)
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(header))}{len(rows)+1}'
    book.save(OUTPUT / 'consolidacao.xlsx')


def main() -> None:
    """Cruza catálogo, plano, fila e exportações somente leitura."""
    OUTPUT.mkdir(parents=True, exist_ok=True)
    status = read_json(ROUND / 'execucao.json')
    closure = read_json(ROUND / 'encerramento.json')
    manifest = read_json(ROUND / 'fila-manifesto.json')
    catalog = load_sites(ROOT / 'var/validacao-final/catalogo-congelado.csv')
    sites = {s.key: s for s in catalog.sites}
    with (ROUND / 'fontes.csv').open(encoding='utf-8-sig', newline='') as stream:
        records = {r['site_key']: r for r in csv.DictReader(stream)}

    # Uma identidade por fonte e URL canônica; bootstrap permanece separado dos XMLs.
    mechanisms: dict[tuple[str, str, bool], dict[str, Any]] = {}

    def mechanism(site: str, url: str, kind: str) -> dict[str, Any]:
        key = site, canonicalize_url(url), kind == 'bootstrap'
        if key not in mechanisms:
            mechanisms[key] = dict(
                fonte=site, nome=sites[site].name, url=url, tipo=kind, no_plano=False,
                raizes_na_fila=0, raizes_fila_concluidas=0, documentos_lidos=0,
                documentos_pendentes=0, xmls_validos=0, documentos_com_falha=0,
                documentos_excluidos=0, ocorrencias_candidatas_nos_xmls=0,
                links_sem_data_nos_xmls=0, filhos_antigos_poupados=0,
                indices_com_limite_profundidade=0, candidatas_distintas_por_raiz=0,
                candidatas_origem_direta=0, exemplos='', diagnostico='',
                categoria='', evidencia='',
            )
        return mechanisms[key]

    plan = read_json(ROUND / 'plano-todos.json')
    for entry in plan:
        for item in entry['mecanismos']:
            mechanism(entry['chave'], item['url'], item['kind'])['no_plano'] = True
    planned_total = len(mechanisms)

    queue = sqlite3.connect(f'file:{ROUND / "fila-documentos.sqlite"}?mode=ro', uri=True)
    queue.row_factory = sqlite3.Row
    roots = {}
    for row in queue.execute('SELECT * FROM roots ORDER BY id'):
        item = mechanism(row['site'], row['url'], row['kind'])
        item['raizes_na_fila'] += 1
        item['raizes_fila_concluidas'] += row['complete']
        roots[row['id']] = item
    document_statuses: Counter[str] = Counter()
    for row in queue.execute('SELECT root_id,state,outcome FROM tasks ORDER BY root_id,id'):
        item = roots[row['root_id']]
        if row['state'] != 'done':
            item['documentos_pendentes'] += 1
            continue
        item['documentos_lidos'] += 1
        outcome = json.loads(row['outcome'])
        document_statuses[outcome['status']] += 1
        for output, field in [
            ('xmls_validos', 'valid'), ('ocorrencias_candidatas_nos_xmls', 'recent'),
            ('links_sem_data_nos_xmls', 'undated'), ('filhos_antigos_poupados', 'old_skipped'),
            ('indices_com_limite_profundidade', 'depth_partial'),
        ]:
            item[output] += outcome.get(field, 0)
        if outcome['status'] in ('blocked', 'error', 'robots_disallowed', 'no_sitemap'):
            item['documentos_com_falha'] += 1
            if not item['diagnostico']:
                item['diagnostico'] = outcome.get('detail') or outcome['status']
        if outcome['status'] == 'excluded':
            item['documentos_excluidos'] += 1
    for row in queue.execute('SELECT root_id,count(*) AS n FROM evidence GROUP BY root_id'):
        roots[row['root_id']]['candidatas_distintas_por_raiz'] += row['n']
    samples: dict[int, list[str]] = defaultdict(list)
    for row in queue.execute('SELECT root_id,url FROM evidence ORDER BY root_id,hash'):
        if len(samples[row['root_id']]) < 3:
            samples[row['root_id']].append(row['url'])
    for root_id, urls in samples.items():
        item = roots[root_id]
        if not item['exemplos']:
            item['exemplos'] = ' | '.join(urls)

    origins: dict[tuple[str, str], dict[str, Any]] = {}
    source_counts: dict[str, Counter[str]] = defaultdict(Counter)
    news_count = 0
    hashes_seen: set[str] = set()
    since = datetime.fromisoformat(status['window_since'])
    until = datetime.fromisoformat(status['window_until'])
    with (ROUND / 'noticias-48h.csv').open(encoding='utf-8-sig', newline='') as stream:
        for row in csv.DictReader(stream):
            news_count += 1
            stamp = datetime.fromisoformat(row['publication_date'] or row['lastmod'])
            assert since <= stamp <= until, 'CSV contém data fora da janela fixa'
            hashes_seen.add(url_hash(row['url']))
            counts = source_counts[row['site_key']]
            counts['total'] += 1
            counts[row['kind']] += 1
            counts[row['base_da_data']] += 1
            key = row['site_key'], canonicalize_url(row['source_sitemap'])
            if key not in origins:
                origins[key] = dict(fonte=key[0], url_origem=row['source_sitemap'],
                                    candidatas=0, publicacao=0, modificacao=0, exemplo=row['url'])
            origin = origins[key]
            origin['candidatas'] += 1
            origin[row['base_da_data']] += 1
    for (site, url), origin in origins.items():
        item = mechanisms.get((site, url, False))
        if item is not None:
            item['candidatas_origem_direta'] += origin['candidatas']
            if not item['exemplos']:
                item['exemplos'] = origin['exemplo']

    source_metrics: dict[str, Counter[str]] = defaultdict(Counter)
    for item in mechanisms.values():
        if item['tipo'] == 'bootstrap':
            category = 'descoberta'
        elif item['documentos_pendentes'] or item['indices_com_limite_profundidade']:
            category = 'parcial'
        elif item['candidatas_distintas_por_raiz'] or item['candidatas_origem_direta']:
            category = 'acessivel_com_candidatas'
        elif item['xmls_validos']:
            category = 'acessivel_sem_candidatas'
        elif item['documentos_com_falha']:
            category = 'falha'
        else:
            category = 'sem_detalhamento'
        item['categoria'] = category
        item['evidencia'] = (
            'fila_por_documento' if item['raizes_na_fila'] else
            'origem_direta_no_csv' if item['candidatas_origem_direta'] else
            'apenas_plano_sem_detalhe_individual'
        )
        metrics = source_metrics[item['fonte']]
        metrics['mecanismos_no_plano'] += int(item['no_plano'])
        metrics['mecanismos_com_evidencia'] += int(
            bool(item['documentos_lidos'] or item['candidatas_origem_direta']))
        for field in ['documentos_lidos', 'documentos_pendentes', 'xmls_validos',
                      'documentos_com_falha', 'indices_com_limite_profundidade']:
            metrics[field] += item[field]

    source_rows: list[dict[str, Any]] = []
    for site_key, site_info in sites.items():
        record = records.get(site_key, {})
        counts = source_counts[site_key]
        category = source_category(site_key in closure['incomplete_sources'],
                                   record.get('status', ''), counts['total'])
        source_rows.append(dict(
            fonte=site_key, nome=site_info.name, pais=site_info.country, url=site_info.url,
            categoria=category, leitura_fonte_concluida=bool(record),
            status_coletor=record.get('status', 'partial'),
            candidatas_exportadas=counts['total'], novas=counts['new'],
            carga_inicial=counts['baseline'], conhecidas=counts['known'],
            data_publicacao=counts['publicacao'], data_modificacao=counts['modificacao'],
            **{field: source_metrics[site_key][field] for field in [
                'mecanismos_no_plano', 'mecanismos_com_evidencia', 'documentos_lidos',
                'documentos_pendentes', 'xmls_validos', 'documentos_com_falha',
                'indices_com_limite_profundidade']},
            total_agregado_coletor=sum(int(record.get(field, 0)) for field in
                                      ['new_urls', 'baseline_urls', 'known_urls']),
            detalhe=closure['reason'] if category == 'parcial' else record.get('detail', ''),
            evidencia_documental=(
                'agregado anterior à fila; contadores de documentos incompletos'
                if site_key in manifest['preserved_sources'] else 'fila por documento'),
        ))

    totals = Counter({field: sum(r[field] for r in source_rows) for field in [
        'candidatas_exportadas', 'novas', 'carga_inicial', 'conhecidas',
        'data_publicacao', 'data_modificacao', 'total_agregado_coletor']})
    total_done = sum(r['documentos_lidos'] for r in source_rows)
    total_pending = sum(r['documentos_pendentes'] for r in source_rows)
    assert len(source_rows) == status['sites_expected'] == 1294
    assert len(records) == status['sites_completed'] == 1293
    assert planned_total == status['mechanisms_selected'] == 38501
    assert total_done == closure['documents_completed']
    assert total_pending == closure['documents_queued']
    assert totals['candidatas_exportadas'] == news_count
    assert totals['data_publicacao'] + totals['data_modificacao'] == news_count
    assert totals['novas'] + totals['carga_inicial'] + totals['conhecidas'] == news_count
    assert news_count - totals['total_agregado_coletor'] == sum(
        source_counts[key]['total'] for key in closure['incomplete_sources'])
    assert all(key in sites for key in records)

    inputs = [ROUND / name for name in ['execucao.json', 'encerramento.json', 'fontes.csv',
                                       'noticias-48h.csv', 'plano-todos.json',
                                       'fila-manifesto.json', 'fila-documentos.sqlite']]
    inputs.append(ROOT / 'var/validacao-final/catalogo-congelado.csv')
    hashes = {}
    for path in inputs:
        with path.open('rb') as binary_stream:
            hashes[str(path.relative_to(ROOT))] = hashlib.file_digest(
                binary_stream, 'sha256').hexdigest()
    summary = dict(
        rodada=12, etapa=1, gerado_em=datetime.now(UTC).isoformat(),
        janela_inicio=status['window_since'], janela_fim=status['window_until'],
        fontes=len(sites), fontes_concluidas=len(records),
        fontes_por_categoria=dict(Counter(r['categoria'] for r in source_rows)),
        fontes_por_status_coletor=dict(Counter(r.get('status', '') for r in records.values())),
        fontes_com_candidatas=sum(r['candidatas_exportadas'] > 0 for r in source_rows),
        fontes_parciais=closure['incomplete_sources'],
        totais=dict(totals), documentos_processados=total_done,
        urls_canonicas_distintas=len(hashes_seen),
        duplicatas_canonicas_no_csv=news_count-len(hashes_seen),
        documentos_pendentes_preservados=total_pending,
        documentos_por_resultado=dict(document_statuses), mecanismos_planejados=planned_total,
        entradas_relatorio_mecanismos=len(mechanisms),
        mecanismos_por_categoria=dict(Counter(r['categoria'] for r in mechanisms.values())),
        origens_diretas_com_candidatas=len(origins),
        fontes_sem_detalhamento_documental_completo=len(manifest['preserved_sources']),
        criterios=CRITERIA,
        entradas_sha256=hashes,
        verificacoes='Totais de fontes, mecanismos, documentos e URLs '
        'reconciliados com as entradas.',
    )
    queue.close()
    source_rows.sort(key=lambda r: r['fonte'])
    mechanism_rows = sorted(mechanisms.values(), key=lambda r: (r['fonte'], r['url'], r['tipo']))
    origin_rows = sorted(origins.values(), key=lambda r: (r['fonte'], r['url_origem']))
    write_csv('fontes.csv', source_rows)
    write_csv('mecanismos.csv', mechanism_rows)
    write_csv('origens-diretas.csv', origin_rows)
    (OUTPUT / 'resumo.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n')
    overview = [dict(indicador=k, valor=v) for k, v in {
        'Fontes do catálogo': len(sites), 'Fontes com leitura encerrada': len(records),
        'Fontes parciais': len(closure['incomplete_sources']),
        'Links candidatos exportados': news_count, 'Novas URLs': totals['novas'],
        'URLs de carga inicial': totals['carga_inicial'],
        'Candidatas com data de publicação': totals['data_publicacao'],
        'Candidatas apenas com data de modificação': totals['data_modificacao'],
        'Tarefas processadas na fila': total_done, 'Tarefas pendentes preservadas': total_pending,
        **{LABELS[k]: v for k, v in summary['fontes_por_categoria'].items()},
    }.items()]
    workbook({'Resumo': overview, 'Fontes': source_rows, 'Mecanismos': mechanism_rows,
              'Origens diretas': origin_rows,
              'Critérios': [dict(criterio=s) for s in CRITERIA]})
    report = [
        '# Rodada 12 — consolidação local da etapa 1', '',
        f"Janela fixa: {status['window_since']} até {status['window_until']}.", '',
        f'{len(records)} fontes com leitura encerrada; NZ Herald parcial. Catálogo: {len(sites)}.',
        '', '| Resultado por fonte | Quantidade |', '|---|---:|',
        *[f'| {LABELS[k]} | {v} |' for k, v in summary['fontes_por_categoria'].items()], '',
        f"O CSV contém **{news_count:,} links candidatos**: {totals['novas']:,} novas URLs e "
        f"{totals['carga_inicial']:,} de carga inicial.", '',
        f"Os agregados por fonte somam {totals['total_agregado_coletor']:,}; a diferença "
        f"de {news_count-totals['total_agregado_coletor']} vem do NZ Herald parcial. "
        'As URLs já estavam gravadas, mas a fonte não recebeu registro final.', '',
        f"**{totals['data_modificacao']:,} candidatas usam somente data de modificação**; "
        f"{totals['data_publicacao']:,} têm metadado de publicação. "
        'A atualidade editorial e o conteúdo textual ainda precisam ser verificados.', '',
        f'{planned_total:,} mecanismos planejados; {len(mechanisms):,} entradas no relatório, '
        'incluindo raízes adicionais e páginas de descoberta geradas durante a execução. '
        f'{len(origins):,} origens diretas de candidatas são exportadas separadamente.', '',
        f'{total_done:,} tarefas processadas e {total_pending:,} pendentes preservadas. '
        'Os totais não equivalem a chamadas HTTP únicas.', '',
        '## Limites de interpretação', '', *[f'- {s}' for s in CRITERIA], '',
        '## Entregas', '',
        '- `consolidacao.xlsx`: resumo, fontes, mecanismos, origens diretas e critérios.',
        '- `fontes.csv`: catálogo completo com resultados e leitura parcial explícita.',
        '- `mecanismos.csv`: inventário planejado e raízes adicionais; evidência por entrada.',
        '- `origens-diretas.csv`: origem gravada nas URLs, inclusive XMLs filhos.',
        '- `resumo.json`: totais reconciliados, critérios e hashes das entradas.', '',
        'Nenhum mecanismo foi excluído ou selecionado. Nenhuma chamada de rede foi feita.', '',
    ]
    (OUTPUT / 'relatorio.md').write_text('\n'.join(report), encoding='utf-8')
    print(json.dumps({k: v for k, v in summary.items()
                      if k not in ('criterios', 'entradas_sha256')}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
