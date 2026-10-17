from pathlib import Path

import openpyxl
import pytest
from tests.strategy.helpers import HOME, NOW, FakeFetcher, robots
from tests.strategy.test_listing import page

from buscador_midias.cli import main
from buscador_midias.state.store import StateStore


def test_comando_mapeia_retoma_e_exporta_sem_escrever_urls_do_corpus(tmp_path: Path):
    sites = tmp_path / "sites.csv"
    sites.write_text(
        "region,name,url,country,country_code,language,source_type\n"
        f"Brasil,Exemplo,{HOME},Brasil,BR,Português,Jornal\n"
    )
    out = tmp_path / "mapa.xlsx"
    args = ["map-sites", "--sites", str(sites), "--pages-per-site", "1", "--out", str(out)]
    fetcher = FakeFetcher(
        {
            HOME + "/robots.txt": robots(),
            HOME: page("/politica"),
            HOME + "/politica": page("/politica/congresso"),
        }
    )
    assert main(args, fetcher=fetcher, now=NOW) == 0
    book = openpyxl.load_workbook(out, read_only=True)
    assert book.sheetnames == [
        "Fontes",
        "Páginas",
        "Mecanismos",
        "Amostras",
        "Cobertura",
        "Lacunas",
        "Complementos",
        "Revisão",
    ]
    assert len(list(book["Páginas"].values)) == 3
    book.close()
    assert main(args, fetcher=fetcher, now=NOW) == 0
    assert fetcher.calls.count(HOME) == 1
    store = StateStore()
    try:
        assert len(store.mapping_pages("exemplo.com")) == 3
        assert store.is_empty()  # nenhum run de descoberta nem URL de notícia
    finally:
        store.close()
    fetcher.calls.clear()
    assert main([*args, "--report-only"], fetcher=fetcher, now=NOW) == 0
    assert fetcher.calls == []


def test_rejeita_orcamento_invalido_antes_de_acessar_rede(tmp_path: Path):
    fetcher = FakeFetcher()
    assert main(["map-sites", "--pages-per-site", "0"], fetcher=fetcher, now=NOW) == 2
    assert fetcher.calls == []


@pytest.mark.parametrize("option,value", [
    ("--max-site-requests", "0"), ("--max-site-requests", "-1"),
    ("--max-site-seconds", "0"), ("--max-site-seconds", "nan"),
    ("--max-site-seconds", "inf"), ("--max-run-seconds", "-1"),
    ("--max-run-seconds", "nan"), ("--max-run-seconds", "inf"),
])
def test_rejeita_limites_totais_invalidos_sem_http(option, value):
    fetcher = FakeFetcher()
    assert main(["map-sites", option, value], fetcher=fetcher, now=NOW) == 2
    assert fetcher.calls == []


def test_prazo_global_impede_http_das_fontes_que_esperaram_na_fila(tmp_path, monkeypatch):
    elapsed = [0]
    monkeypatch.setattr("buscador_midias.cli.monotonic", lambda: elapsed[0])

    class SlowFetcher(FakeFetcher):
        def get(self, url):
            elapsed[0] = 2
            return super().get(url)

    sites = tmp_path / "sites.csv"
    sites.write_text(
        "region,name,url,country,country_code,language,source_type\n"
        f"Brasil,Exemplo,{HOME},Brasil,BR,Português,Jornal\n"
        "Brasil,Outro,https://outro.com,Brasil,BR,Português,Jornal\n"
    )
    fetcher = SlowFetcher({HOME + "/robots.txt": robots()})
    assert main([
        "map-sites", "--sites", str(sites), "--drain", "--workers", "1",
        "--max-run-seconds", "1", "--out", str(tmp_path / "mapa.xlsx"),
    ], fetcher=fetcher, now=NOW) == 0
    assert fetcher.calls == [HOME + "/robots.txt"]
    store = StateStore()
    try:
        for key in ("exemplo.com", "outro.com"):
            root = store.mapping_pages(key)[0]
            assert root.status.value == "pending"
            assert "limite de tempo" in root.detail
    finally:
        store.close()


def test_prioridades_exige_modo_somente_leitura_e_nao_faz_http(tmp_path: Path):
    fetcher = FakeFetcher()
    assert main(["map-sites", "--priorities-only"], fetcher=fetcher, now=NOW) == 2
    sites = tmp_path / "sites.csv"
    sites.write_text(
        "region,name,url,country,country_code,language,source_type\n"
        f"Brasil,Exemplo,{HOME},Brasil,BR,Português,Jornal\n"
    )
    out = tmp_path / "prioridades.xlsx"
    assert main([
        "map-sites", "--sites", str(sites), "--report-only", "--priorities-only",
        "--out", str(out),
    ], fetcher=fetcher, now=NOW) == 0
    book = openpyxl.load_workbook(out, read_only=True)
    assert book.sheetnames == ["Prioridades"]
    assert list(book["Prioridades"].values)[1][1] == "exemplo.com"
    book.close()
    assert fetcher.calls == []


def test_drain_termina_com_urls_sempre_novas_e_preserva_fila(tmp_path: Path):
    from buscador_midias.net.fetcher import FetchResult

    class GrowingFetcher(FakeFetcher):
        def get(self, url):
            self.calls.append(url)
            body = robots() if url.endswith("/robots.txt") else page(
                "/" + "a" * (len(url.removeprefix(HOME).strip("/")) + 1)
            )
            return FetchResult(url, 200, body, None)

    sites = tmp_path / "sites.csv"
    sites.write_text(
        "region,name,url,country,country_code,language,source_type\n"
        f"Brasil,Exemplo,{HOME},Brasil,BR,Português,Jornal\n"
    )
    fetcher = GrowingFetcher()
    assert main([
        "map-sites", "--sites", str(sites), "--drain", "--pages-per-site", "2",
        "--max-site-requests", "5", "--out", str(tmp_path / "mapa.xlsx"),
    ], fetcher=fetcher, now=NOW) == 0
    assert len(fetcher.calls) == 5
    store = StateStore()
    try:
        pages = store.mapping_pages("exemplo.com")
        assert any(p.status.value == "pending" for p in pages)
        root = next(p for p in pages if p.kind == "root")
        assert root.status.value == "ok"
        assert "limite total de requisições" in root.detail
        assert not any(p.status.value == "error" for p in pages)
        assert store.is_empty()
    finally:
        store.close()


def test_orcamento_total_tambem_limita_auditoria_e_preserva_checkpoints(tmp_path: Path):
    from tests.strategy.helpers import rss

    from buscador_midias.domain.mapping import MapPage, PageStatus

    sites = tmp_path / "sites.csv"
    sites.write_text(
        "region,name,url,country,country_code,language,source_type\n"
        f"Brasil,Exemplo,{HOME},Brasil,BR,Português,Jornal\n"
    )
    feeds = tuple(HOME + f"/feed-{i}.xml" for i in range(8))
    store = StateStore()
    try:
        store.save_mapping_pages("exemplo.com", [
            MapPage(HOME, "", "root", PageStatus.OK, feeds=feeds, checked_at=NOW),
        ])
    finally:
        store.close()
    fetcher = FakeFetcher({HOME + "/robots.txt": robots(), **{u: rss() for u in feeds}})
    assert main([
        "map-sites", "--sites", str(sites), "--drain", "--sources-only",
        "--max-site-requests", "5", "--out", str(tmp_path / "mapa.xlsx"),
    ], fetcher=fetcher, now=NOW) == 0
    assert len(fetcher.calls) == 5
    store = StateStore()
    try:
        assert len(store.mapping_sources("exemplo.com")) == 4
        assert "limite total de requisições" in store.mapping_pages("exemplo.com")[0].detail
    finally:
        store.close()


def test_drain_does_not_restart_paused_source_audits_in_the_same_run(tmp_path: Path):
    from buscador_midias.net.fetcher import FetchResult

    sites = tmp_path / "sites.csv"
    sites.write_text(
        "region,name,url,country,country_code,language,source_type\n"
        f"Brasil,Exemplo,{HOME},Brasil,BR,Português,Jornal\n"
    )
    urls = [HOME + f"/feed-{i}.xml" for i in range(5)]
    markup = "<html><head>" + "".join(
        f'<link rel="alternate" type="application/rss+xml" href="{url}">'
        for url in urls
    ) + "</head><body></body></html>"
    fetcher = FakeFetcher({HOME + "/robots.txt": robots(), HOME: markup.encode()})
    fetcher.responses.update({
        url: FetchResult(url, None, b"", "ConnectError: Temporary failure in name resolution")
        for url in urls
    })
    assert main([
        "map-sites", "--sites", str(sites), "--drain", "--validate-sources",
        "--out", str(tmp_path / "mapa.xlsx"),
    ], fetcher=fetcher, now=NOW) == 0
    assert not set(urls[3:]).intersection(fetcher.calls)
    store = StateStore()
    try:
        assert len(store.mapping_sources("exemplo.com")) == 3
        assert store.is_empty()
    finally:
        store.close()


def test_drain_mapeia_prefixos_novos_de_outro_host_apos_suspender_subdominio(tmp_path: Path):
    from tests.strategy.helpers import plain_sitemap

    from buscador_midias.net.fetcher import FetchResult

    sites = tmp_path / "sites.csv"
    sites.write_text(
        "region,name,url,country,country_code,language,source_type\n"
        f"Brasil,Exemplo,{HOME},Brasil,BR,Português,Jornal\n"
    )
    bad = [f"https://falha.exemplo.com/secao-{c}" for c in "abcde"]
    feed = HOME + "/feed.xml"
    fetcher = FakeFetcher({
        HOME + "/robots.txt": robots(), "https://falha.exemplo.com/robots.txt": robots(),
        HOME: page(*bad, "/politica"),
        HOME + "/politica": b'<html><link rel="alternate" type="application/rss+xml" '
                            b'href="/feed.xml"></html>',
        HOME + "/oculta": page("/oculta/materia-com-titulo-do-jornal"),
        feed: plain_sitemap(("oculta/materia-com-titulo-do-jornal", 1)),
        **{u: FetchResult(u, None, b"", "ConnectError: DNS indisponivel") for u in bad},
    })
    assert main([
        "map-sites", "--sites", str(sites), "--drain", "--validate-sources",
        "--pages-per-site", "4", "--out", str(tmp_path / "mapa.xlsx"),
    ], fetcher=fetcher, now=NOW) == 0
    assert HOME + "/oculta" in fetcher.calls
    assert not set(bad[3:]).intersection(fetcher.calls)
    assert fetcher.calls.count(feed) == 1
    store = StateStore()
    try:
        pending = {p.url for p in store.mapping_pages("exemplo.com") if p.status.value == "pending"}
        assert pending == set(bad[3:])
        assert store.is_empty()
    finally:
        store.close()


def test_drain_completa_subsecoes_e_mecanismos_em_uma_execucao(tmp_path: Path):
    from tests.strategy.helpers import plain_sitemap

    sites = tmp_path / "sites.csv"
    sites.write_text(
        "region,name,url,country,country_code,language,source_type\n"
        f"Brasil,Exemplo,{HOME},Brasil,BR,Português,Jornal\n"
    )
    feed = HOME + "/feed.xml"
    slug = "/politica/materia-com-titulo-do-jornal"
    fetcher = FakeFetcher(
        {
            HOME + "/robots.txt": robots(),
            HOME: page("/politica"),
            HOME + "/politica": (
                '<html><head><link rel="alternate" type="application/rss+xml" href="/feed.xml">'
                f'</head><body><a href="{slug}">Noticia</a></body></html>'
            ).encode(),
            HOME + "/oculta": page("/oculta/materia-com-titulo-do-jornal"),
            feed: plain_sitemap((slug.lstrip("/"), 1), ("oculta/materia-com-titulo-do-jornal", 1)),
        }
    )
    args = [
        "map-sites",
        "--sites",
        str(sites),
        "--drain",
        "--pages-per-site",
        "1",
        "--validate-sources",
        "--source-requests",
        "2",
        "--out",
        str(tmp_path / "mapa.xlsx"),
    ]
    assert main(args, fetcher=fetcher, now=NOW) == 0
    store = StateStore()
    try:
        assert all(p.status.value != "pending" for p in store.mapping_pages("exemplo.com"))
        assert store.mapping_sources("exemplo.com")[0].checked_at == NOW
        derived = next(p for p in store.mapping_pages("exemplo.com") if p.url == HOME + "/oculta")
        assert derived.found_on == feed and derived.discovery_basis == "mechanism_path"
        assert store.is_empty()
    finally:
        store.close()
    assert fetcher.calls.count(feed) == 1
    assert HOME + slug not in fetcher.calls
    assert HOME + "/oculta/materia-com-titulo-do-jornal" not in fetcher.calls


def test_confere_e_persiste_mecanismo_sem_baixar_materia(tmp_path: Path):
    from tests.strategy.helpers import plain_sitemap

    sites = tmp_path / "sites.csv"
    sites.write_text(
        "region,name,url,country,country_code,language,source_type\n"
        f"Brasil,Exemplo,{HOME},Brasil,BR,Português,Jornal\n"
    )
    slug = "/economia/materia-com-titulo-do-jornal"
    fetcher = FakeFetcher(
        {
            HOME + "/robots.txt": robots(),
            HOME: b'<html><link rel="alternate" type="application/rss+xml" href="/feed.xml">'
            + page(slug)
            + b"</html>",
            HOME + "/feed.xml": plain_sitemap((slug.lstrip("/"), 1)),
        }
    )
    out = tmp_path / "mapa.xlsx"
    args = [
        "map-sites",
        "--sites",
        str(sites),
        "--pages-per-site",
        "1",
        "--validate-sources",
        "--hints",
        "/nao/existe.csv",
        "--out",
        str(out),
    ]
    assert main(args, fetcher=fetcher, now=NOW) == 0
    store = StateStore()
    try:
        sources = store.mapping_sources("exemplo.com")
        assert len(sources) == 1
        assert sources[0].listed_urls == (HOME + slug,)
        assert sources[0].checked_at == NOW
        assert store.is_empty()
    finally:
        store.close()
    assert HOME + slug not in fetcher.calls
    fetcher.calls.clear()
    assert main([*args, "--sources-only"], fetcher=fetcher, now=NOW) == 0
    assert fetcher.calls == []  # a leitura persistida já atende a amostra
    book = openpyxl.load_workbook(out, read_only=True)
    assert list(book["Cobertura"].values)[1][2] == "sample_full"
    book.close()


def test_auditoria_retoma_orcamento_no_proximo_lote_sem_esperar_24_h(tmp_path: Path):
    from tests.strategy.helpers import plain_sitemap, sitemap_index

    from buscador_midias.domain.mapping import MapPage, PageStatus

    sites = tmp_path / "sites.csv"
    sites.write_text(
        "region,name,url,country,country_code,language,source_type\n"
        f"Brasil,Exemplo,{HOME},Brasil,BR,Português,Jornal\n"
    )
    index, child = HOME + "/index.xml", HOME + "/child.xml"
    article = HOME + "/politica/materia-com-titulo-do-jornal"
    store = StateStore()
    try:
        store.save_mapping_pages(
            "exemplo.com",
            [
                MapPage(
                    HOME,
                    "",
                    "root",
                    PageStatus.OK,
                    article_urls=(article,),
                    feeds=(index,),
                    checked_at=NOW,
                )
            ],
        )
    finally:
        store.close()
    fetcher = FakeFetcher(
        {
            HOME + "/robots.txt": robots(),
            index: sitemap_index((child, 1)),
            child: plain_sitemap(("politica/materia-com-titulo-do-jornal", 1)),
        }
    )
    args = [
        "map-sites",
        "--sites",
        str(sites),
        "--sources-only",
        "--hints",
        "/nao/existe.csv",
        "--out",
        str(tmp_path / "mapa.xlsx"),
    ]
    assert main([*args, "--source-requests", "2"], fetcher=fetcher, now=NOW) == 0
    assert child not in fetcher.calls
    assert main([*args, "--source-requests", "3"], fetcher=fetcher, now=NOW) == 0
    assert child in fetcher.calls
    store = StateStore()
    try:
        assert store.mapping_sources("exemplo.com")[0].listed_urls == (article,)
    finally:
        store.close()
