import csv

import openpyxl
from tests.strategy.helpers import HOME, NOW, make_site

from buscador_midias.domain.mapping import MapPage, MapSource, PageStatus
from buscador_midias.domain.strategy import (
    CandidateOrigin,
    CandidateReport,
    StrategyKind,
    Verdict,
)
from buscador_midias.strategy.mapping_export import mapping_tables, write_mapping


def test_planilha_divide_abas_no_limite_do_excel_sem_perder_linhas(tmp_path, monkeypatch):
    import buscador_midias.strategy.mapping_export as exporting

    monkeypatch.setattr(exporting, "MAX_XLSX_ROWS", 4)
    out = tmp_path / "mapa.xlsx"
    write_mapping({"Amostras": (("URL",), [(str(i),) for i in range(8)])}, out)
    book = openpyxl.load_workbook(out, read_only=True)
    assert book.sheetnames == ["Amostras", "Amostras 2", "Amostras 3"]
    found = []
    for sheet in book:
        rows = list(sheet.values)
        assert rows[0] == ("URL",) and len(rows) <= 4
        found.extend(row[0] for row in rows[1:])
    assert found == [str(i) for i in range(8)]
    book.close()


def test_relatorio_nao_declara_cobertura_so_porque_a_fila_acabou():
    pages = [MapPage(HOME, "", "root", PageStatus.OK)]
    tables = mapping_tables([make_site()], {"exemplo.com": pages}, {})
    row = dict(zip(tables["Fontes"][0], tables["Fontes"][1][0], strict=True))
    assert row["Exploração"] == "fila explorada"
    assert row["Cobertura de notícias"] == "não validada"


def test_reaproveita_candidato_anterior_sem_inventar_a_secao_que_ele_cobre():
    report = CandidateReport(
        url=HOME + "/rss.xml",
        origin=CandidateOrigin.HINT,
        verdict=Verdict.FRESH,
        kind=StrategyKind.FEED,
        measure=None,
        score=1,
        detail="",
    )
    tables = mapping_tables([make_site()], {}, {}, [("exemplo.com", report, NOW)])
    row = dict(zip(tables["Mecanismos"][0], tables["Mecanismos"][1][0], strict=True))
    assert row["Página"] == ""
    assert row["Mecanismo"] == report.url
    assert row["Validação"] == "fresh; cobertura por seção pendente"


def test_todos_os_mecanismos_e_amostras_sao_exportados_com_protecao_de_formula(tmp_path):
    pages = [
        MapPage(
            HOME,
            "",
            "root",
            PageStatus.OK,
            title="=1+1",
            article_urls=tuple(HOME + f"/artigo-{i}" for i in range(70)),
            feeds=(HOME + "/rss.xml",),
            feed_indexes=(HOME + "/rss.html",),
        )
    ]
    tables = mapping_tables([make_site()], {"exemplo.com": pages}, {})
    assert len(tables["Amostras"][1]) == 70
    assert len(tables["Mecanismos"][1]) == 2
    out = tmp_path / "mapa.xlsx"
    write_mapping(tables, out)
    book = openpyxl.load_workbook(out)
    assert book["Páginas"]["G2"].value == "'=1+1"
    assert book["Amostras"].max_row == 71
    book.close()
    write_mapping(tables, tmp_path / "mapa.csv")
    with (tmp_path / "mapa-Páginas.csv").open(encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    assert rows[1][6] == "'=1+1"


def test_titulo_com_caractere_de_controle_nao_impede_exportacao_do_catalogo(tmp_path):
    tables = mapping_tables(
        [make_site()],
        {
            "exemplo.com": [
                MapPage(HOME, "", "root", PageStatus.OK, title="Notícia\x07 válida"),
            ]
        },
        {},
    )
    out = tmp_path / "mapa.xlsx"
    write_mapping(tables, out)
    book = openpyxl.load_workbook(out, read_only=True)
    assert list(book["Páginas"].values)[1][6] == "Notícia válida"
    book.close()


def test_resumo_e_lacunas_mostram_correspondencia_parcial_sem_afirmar_cobertura_total():
    first, missing, source = HOME + "/primeira", HOME + "/faltante", HOME + "/feed.xml"
    page = MapPage(HOME, "", "root", PageStatus.OK, article_urls=(first, missing), checked_at=NOW)
    checks = [MapSource(source, "fresh", "feed", (first,), NOW)]
    tables = mapping_tables(
        [make_site()], {"exemplo.com": [page]}, {}, checks={"exemplo.com": checks}
    )
    summary = dict(zip(tables["Fontes"][0], tables["Fontes"][1][0], strict=True))
    assert summary["Amostras parciais"] == 1
    assert summary["Amostras cobertas"] == 0
    assert summary["Cobertura de notícias"] == "não validada"
    assert tables["Lacunas"][1] == [("exemplo.com", HOME, "partial", missing)]
    suggestion = dict(zip(tables["Complementos"][0], tables["Complementos"][1][0], strict=True))
    assert suggestion["Situação da amostra"] == "partial"
    assert suggestion["Mecanismo"] == source
    assert suggestion["Links acrescentados"] == 1
    assert suggestion["Links correspondentes"] == 1
    assert suggestion["Leitura parcial"] == "não"
    assert suggestion["Mecanismo lido em (UTC)"] == NOW.isoformat()


def test_materia_identificada_nao_e_contada_como_secao_ou_lacuna_de_cobertura():
    classified = MapPage(
        HOME + "/titulo",
        HOME,
        "section",
        PageStatus.EXCLUDED,
        page_type="article",
        type_basis="self_article_schema",
        discovery_basis="navigation",
    )
    tables = mapping_tables([make_site()], {"exemplo.com": [classified]}, {})
    row = dict(zip(tables["Fontes"][0], tables["Fontes"][1][0], strict=True))
    assert row["Seções candidatas"] == 0
    assert row["Matérias identificadas"] == 1
    assert tables["Cobertura"][1] == []
    inventory = dict(zip(tables["Páginas"][0], tables["Páginas"][1][0], strict=True))
    assert inventory["Tipo identificado"] == "article"
    assert inventory["Evidência do tipo"] == "self_article_schema"
