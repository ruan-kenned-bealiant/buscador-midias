"""Planilha de conferência: o que se sabe de cada fonte, os candidatos e as seções."""

from dataclasses import replace
from datetime import UTC, datetime, timedelta
from pathlib import Path

from buscador_midias.domain.site import Site
from buscador_midias.domain.strategy import (
    CandidateOrigin,
    CandidateReport,
    Measure,
    ProbeState,
    Role,
    StrategyKind,
    Verdict,
)
from buscador_midias.strategy.export import build_tables, write_csv, write_xlsx
from buscador_midias.strategy.hints import SourceHints

T0 = datetime(2026, 10, 16, 12, 0, tzinfo=UTC)


def site(name: str, url: str) -> Site:
    return Site.create(
        name=name,
        url=url,
        region="Brasil",
        country="Brasil",
        country_code="BR",
        language="Português",
        source_type="Portal",
    )


def candidate(
    url: str,
    role: Role | None,
    *,
    kind=StrategyKind.NEWS_SITEMAP,
    verdict=Verdict.FRESH,
    recent=100,
    coverage=0.8,
    sections=(("sp", 60), ("politica", 40)),
    value=100.0,
) -> CandidateReport:
    measure = Measure(
        kind, 120, 120, recent, 120, T0, 2048, sections=len(sections), top_sections=sections
    )
    return CandidateReport(
        url=url,
        origin=CandidateOrigin.ROBOTS,
        verdict=verdict,
        kind=kind,
        measure=measure,
        score=value,
        detail="",
        requests=3,
        role=role,
        coverage=coverage,
        union_recent=125,
        cover_share=0.96,
    )


G1 = site("g1", "https://g1.globo.com")
BLOQUEADA = site("AP", "https://apnews.com")
NOVA = site("Nova", "https://nova.com")
ROWS = [
    ("g1.globo.com", candidate("https://g1.globo.com/sitemap.xml", Role.PRIMARY), T0),
    (
        "g1.globo.com",
        candidate(
            "https://g1.globo.com/rss/g1/",
            Role.COMPLEMENT,
            kind=StrategyKind.FEED,
            recent=25,
            coverage=0.2,
            sections=(("mg", 25),),
        ),
        T0,
    ),
    (
        "g1.globo.com",
        candidate(
            "https://g1.globo.com/rss/velho", None, verdict=Verdict.STALE, recent=0, coverage=None
        ),
        T0,
    ),
]
STATES = {
    "apnews.com": ProbeState(
        "apnews.com", T0, False, 2, T0 + timedelta(hours=12), T0 + timedelta(hours=12), "blocked: x"
    ),
    "g1.globo.com": ProbeState("g1.globo.com", T0, True, 0, T0, T0 + timedelta(days=7)),
}
HINTS = SourceHints(
    {}, {"g1.globo.com": {"categoria": "C - Sitemap genérico", "gdelt_24h": "2298"}}
)


def tables():
    return build_tables([G1, BLOQUEADA, NOVA], ROWS, STATES, HINTS)


def summary_row(key: str) -> dict[str, object]:
    t = tables().summary
    row = next(r for r in t.rows if r[t.headers.index("Chave")] == key)
    return dict(zip(t.headers, row, strict=True))


def test_uma_linha_por_fonte_do_catalogo_na_ordem_dada():
    t = tables().summary

    assert [r[t.headers.index("Chave")] for r in t.rows] == [
        "g1.globo.com",
        "apnews.com",
        "nova.com",
    ]


def test_fonte_com_estrategia_mostra_principal_complemento_e_cobertura():
    row = summary_row("g1.globo.com")

    assert row["Situação"] == "ok"
    assert row["Estratégia principal"] == "news_sitemap"
    assert row["Endereço principal"] == "https://g1.globo.com/sitemap.xml"
    assert row["Recentes 48h (principal)"] == 100
    assert row["Seções (principal)"] == 2
    assert row["Maior seção"] == "sp (60%)"
    assert row["Complementos"] == 1
    assert row["Endereços dos complementos"] == "https://g1.globo.com/rss/g1/"
    assert row["Cobertura do principal"] == 0.8
    assert row["Cobertura do conjunto"] == 0.96
    assert row["Matérias recentes (união 48h)"] == 125
    assert row["Estimativa por dia"] == 62  # 125 em 48 h, arredondado
    assert row["Ref.: GDELT/dia"] == "2298"
    assert row["Ref.: categoria"] == "C - Sitemap genérico"


def test_fonte_sem_estrategia_traz_o_motivo_e_a_espera():
    row = summary_row("apnews.com")

    assert row["Situação"] == "sem estratégia"
    assert row["Motivo (sem estratégia)"] == "blocked: x"
    assert row["Tentativas sem melhora"] == 2


def test_fonte_nunca_sondada_aparece_como_nao_sondada():
    row = summary_row("nova.com")

    assert row["Situação"] == "não sondada"
    assert row["Endereço principal"] == ""


def test_candidatos_vao_para_outra_aba_com_papel_veredito_e_cobertura():
    t = tables().candidates
    rows = [dict(zip(t.headers, r, strict=True)) for r in t.rows]

    assert [r["Papel"] for r in rows] == ["principal", "complemento", ""]
    assert rows[2]["Veredito"] == "stale"
    assert rows[1]["Cobertura"] == 0.2
    assert all(r["Fonte"] == "g1" for r in rows)


def test_secoes_dos_candidatos_escolhidos_com_a_parte_de_cada_uma():
    t = tables().sections
    rows = [dict(zip(t.headers, r, strict=True)) for r in t.rows]

    assert [(r["Seção"], r["Matérias recentes"], r["% do candidato"]) for r in rows] == [
        ("sp", 60, 0.6),
        ("politica", 40, 0.4),
        ("mg", 25, 1.0),
    ]
    assert rows[2]["Papel"] == "complemento"


def test_xlsx_tem_tres_abas_com_cabecalho_e_dados(tmp_path: Path):
    from openpyxl import load_workbook

    path = tmp_path / "estrategias.xlsx"

    write_xlsx(tables(), path)

    book = load_workbook(path)
    assert book.sheetnames == ["Estratégia", "Candidatos", "Seções"]
    sheet = book["Estratégia"]
    assert sheet.cell(1, 1).value == "Fonte"
    assert sheet.max_row == 4  # cabeçalho + 3 fontes
    assert sheet.freeze_panes == "C2"
    assert sheet.auto_filter.ref


def test_csv_grava_um_arquivo_por_aba_ao_lado(tmp_path: Path):
    path = tmp_path / "estrategias.csv"

    write_csv(tables(), path)

    assert path.read_text(encoding="utf-8-sig").splitlines()[0].startswith("Fonte,")
    assert (tmp_path / "estrategias-candidatos.csv").is_file()
    assert (tmp_path / "estrategias-secoes.csv").is_file()


def test_fonte_lenta_aparece_como_lenta_com_a_data_da_ultima_materia():
    lenta = site("Lenta", "https://lenta.com")
    rows = [
        (
            "lenta.com",
            candidate(
                "https://lenta.com/feed/",
                Role.PRIMARY,
                kind=StrategyKind.FEED,
                verdict=Verdict.SLOW,
                recent=0,
                coverage=None,
            ),
            T0,
        )
    ]

    t = build_tables([lenta], rows, {}, SourceHints({}, {})).summary
    row = dict(zip(t.headers, t.rows[0], strict=True))

    assert row["Situação"] == "lenta"
    assert row["Recentes 48h (principal)"] == 0
    assert row["Última matéria (principal)"] == T0


def _map(rows, sites=None):
    t = build_tables(sites or [G1], rows, {}, SourceHints({}, {})).section_map
    return t, [dict(zip(t.headers, r, strict=True)) for r in t.rows]


def test_mapa_de_secoes_junta_as_secoes_de_todas_as_candidatas_e_marca_as_cobertas():
    rows = [
        ("g1.globo.com", candidate("https://g1.globo.com/a.xml", Role.PRIMARY), T0),
        (
            "g1.globo.com",
            candidate(
                "https://g1.globo.com/b.xml", Role.COMPLEMENT, sections=(("mg", 25), ("sp", 5))
            ),
            T0,
        ),
        (
            "g1.globo.com",
            candidate("https://g1.globo.com/c.xml", None, sections=(("rj", 30),), recent=30),
            T0,
        ),
    ]

    _, mapped = _map(rows)
    by_section = {r["Seção"]: r for r in mapped}

    assert set(by_section) == {"sp", "politica", "mg", "rj"}
    assert by_section["sp"]["Coberta pelo conjunto"] == "sim"
    assert by_section["sp"]["Matérias"] == 60  # a maior contagem entre as candidatas
    assert by_section["mg"]["Listada em"] == "complemento"
    assert by_section["rj"]["Coberta pelo conjunto"] == "não"
    assert by_section["rj"]["Listada em"] == "só em candidata não escolhida"
    assert [r["Seção"] for r in mapped][0] == "sp"  # da maior para a menor


def test_resumo_conta_secoes_vistas_e_fora_do_conjunto():
    rows = [
        ("g1.globo.com", candidate("https://g1.globo.com/a.xml", Role.PRIMARY), T0),
        (
            "g1.globo.com",
            candidate("https://g1.globo.com/c.xml", None, sections=(("rj", 30),), recent=30),
            T0,
        ),
    ]

    t = build_tables([G1], rows, {}, SourceHints({}, {})).summary
    row = dict(zip(t.headers, t.rows[0], strict=True))

    assert row["Seções vistas (todas as candidatas)"] == 3
    assert row["Seções fora do conjunto"] == 1


def test_url_plana_nao_tem_secoes_confiaveis():
    """Quando o primeiro trecho do caminho muda a cada matéria, não é seção."""
    flat = tuple((f"titulo-{n}", 1) for n in range(40))
    rows = [
        (
            "g1.globo.com",
            candidate("https://g1.globo.com/a.xml", Role.PRIMARY, recent=40, sections=flat),
            T0,
        )
    ]

    t = build_tables([G1], rows, {}, SourceHints({}, {}))
    mapped = [dict(zip(t.section_map.headers, r, strict=True)) for r in t.section_map.rows]
    row = dict(zip(t.summary.headers, t.summary.rows[0], strict=True))

    assert len(mapped) == 1 and "não são seções" in mapped[0]["Observação"]
    assert row["Seções vistas (todas as candidatas)"] == "URL plana"


def test_secoes_alem_do_limite_guardado_viram_um_aviso():
    secs = tuple((f"s{n}", 100 - n) for n in range(40))
    base = candidate("https://g1.globo.com/a.xml", Role.PRIMARY, recent=2000, sections=secs)
    assert base.measure is not None
    truncated = replace(base, measure=replace(base.measure, sections=90))
    rows = [("g1.globo.com", truncated, T0)]

    _, mapped = _map(rows)

    assert any("90 seções" in r["Observação"] for r in mapped)


def test_arquivos_csv_ganham_o_mapa_de_secoes(tmp_path: Path):
    tables = build_tables([G1], ROWS, STATES, HINTS)

    write_csv(tables, tmp_path / "e.csv")

    assert (tmp_path / "e-mapa-secoes.csv").is_file()
