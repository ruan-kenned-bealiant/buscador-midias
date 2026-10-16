from pathlib import Path

import pytest
from tests.strategy.helpers import NOW, FakeFetcher, news_sitemap, robots, rss, status

from buscador_midias.cli import build_parser, main

SITES = (
    "region,name,url,country,country_code,language,source_type\n"
    "Brasil,A,https://a.com,Brasil,BR,Português,Jornal\n"
    "Brasil,B,https://b.com,Brasil,BR,Português,Jornal\n"
    "Brasil,C,https://c.com,Brasil,BR,Português,Jornal\n"
)


@pytest.fixture
def sites(tmp_path: Path) -> Path:
    path = tmp_path / "sites.csv"
    path.write_text(SITES, encoding="utf-8")
    return path


def probe(sites: Path) -> None:
    fetcher = FakeFetcher(
        {
            "https://a.com/robots.txt": robots("https://a.com/news.xml"),
            "https://a.com/news.xml": news_sitemap(
                *((f"politica/n{i}", 1) for i in range(12)), host="https://a.com"
            ),
            "https://a.com": b'<link rel="alternate" type="application/rss+xml" href="/feed">',
            "https://a.com/feed": rss(
                *((f"mundo/f{i}", 1) for i in range(8)), host="https://a.com"
            ),
            "https://b.com/robots.txt": status("https://b.com/robots.txt", 403),
        }
    )
    code = main(
        [
            "probe-sources",
            "--sites",
            str(sites),
            "--hints",
            "/nao/existe.csv",
            "--only",
            "a.com",
            "b.com",
        ],
        fetcher=fetcher,
        now=NOW,
    )
    assert code == 0


def test_parser_do_export_tem_saida_padrao_em_xlsx():
    args = build_parser().parse_args(["export-strategies"])

    assert args.out.name == "estrategias.xlsx"


def test_exporta_todas_as_fontes_do_catalogo_inclusive_as_nao_sondadas(
    sites: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
):
    from openpyxl import load_workbook

    probe(sites)
    out = tmp_path / "saida.xlsx"

    code = main(
        [
            "export-strategies",
            "--sites",
            str(sites),
            "--hints",
            "/nao/existe.csv",
            "--out",
            str(out),
        ]
    )

    assert code == 0 and out.is_file()
    sheet = load_workbook(out)["Estratégia"]
    header = [c.value for c in sheet[1]]
    situation = {
        sheet.cell(r, header.index("Chave") + 1).value: sheet.cell(
            r, header.index("Situação") + 1
        ).value
        for r in range(2, sheet.max_row + 1)
    }
    assert situation == {"a.com": "ok", "b.com": "sem estratégia", "c.com": "não sondada"}
    assert "3 fontes" in capsys.readouterr().out


def test_exporta_csv_pelo_sufixo_e_so_as_fontes_pedidas(sites: Path, tmp_path: Path):
    probe(sites)
    out = tmp_path / "saida.csv"

    main(
        [
            "export-strategies", "--sites", str(sites), "--hints", "/nao/existe.csv",
            "--only", "a.com", "--out", str(out),
        ]
    )  # fmt: skip

    lines = out.read_text(encoding="utf-8-sig").splitlines()
    assert len(lines) == 2 and lines[1].startswith("A,a.com")
    assert (tmp_path / "saida-candidatos.csv").is_file()


def test_sufixo_desconhecido_e_erro_de_uso(
    sites: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
):
    code = main(["export-strategies", "--sites", str(sites), "--out", str(tmp_path / "x.txt")])

    assert code == 2
    assert ".xlsx" in capsys.readouterr().err
