"""Relatório do inventário completo, sem confundir exploração com cobertura."""

from __future__ import annotations

import csv
import re
from collections import Counter
from collections.abc import Iterator, Mapping, Sequence
from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile

from buscador_midias.domain.mapping import MapPage, MapSource
from buscador_midias.domain.site import Site
from buscador_midias.domain.strategy import CandidateReport
from buscador_midias.strategy.mapping import summarize_map
from buscador_midias.strategy.mapping_coverage import CoverageIndex
from buscador_midias.strategy.mapping_review import review_entry

type Cell = str | int | float | None
type Sheet = tuple[tuple[str, ...], list[tuple[Cell, ...]]]
MAX_XLSX_ROWS = 1_048_576


def _sheet_parts(tables: Mapping[str, Sheet]) -> Iterator[tuple[str, Sheet]]:
    """Divide abas grandes preservando cabeçalho e todas as linhas."""
    size = MAX_XLSX_ROWS - 1
    for name, (headers, rows) in tables.items():
        for part, offset in enumerate(range(0, max(1, len(rows)), size), 1):
            suffix = f" {part}" if part > 1 else ""
            yield name[: 31 - len(suffix)] + suffix, (headers, rows[offset : offset + size])


def mapping_tables(
    sites: Sequence[Site],
    pages: Mapping[str, Sequence[MapPage]],
    previous: Mapping[str, CandidateReport],
    candidates: Sequence[tuple[str, CandidateReport, datetime]] = (),
    checks: Mapping[str, Sequence[MapSource]] | None = None,
) -> dict[str, Sheet]:
    """Monta resumo, páginas, mecanismos declarados e links de amostra, sem truncar."""
    summary: Sheet = (
        (
            "Fonte",
            "Chave",
            "País",
            "Exploração",
            "Seções candidatas",
            "Páginas lidas",
            "Pendentes",
            "Falhas/proibições",
            "Páginas excluídas",
            "Mecanismos encontrados",
            "Cobertura de notícias",
            "Principal anterior",
            "Observação",
            "Páginas com amostra",
            "Amostras cobertas",
            "Amostras parciais",
            "Amostras sem correspondência",
            "Páginas lidas sem amostra",
            "Mecanismos conferidos",
            "Matérias identificadas",
            "Serviços identificados",
            "Seções confirmadas por metadados",
            "Candidatas apontadas na navegação",
            "Candidatas derivadas de mecanismos",
        ),
        [],
    )
    inventory: Sheet = (
        (
            "Chave",
            "Fonte",
            "Página",
            "Encontrada em",
            "Tipo",
            "Estado",
            "Título",
            "Links de matéria",
            "Tentativas",
            "Verificada em (UTC)",
            "Detalhe",
            "Cobertura de notícias",
            "Tipo identificado",
            "Evidência do tipo",
            "Como foi encontrada",
        ),
        [],
    )
    sources: Sheet = (
        (
            "Chave",
            "Página",
            "Mecanismo",
            "Tipo de pista",
            "Validação",
            "Verificado em (UTC)",
            "URLs listadas",
            "Leitura parcial",
            "Detalhe",
        ),
        [],
    )
    samples: Sheet = (("Chave", "Página", "Link de matéria (heurística)"), [])
    coverage: Sheet = (
        (
            "Chave",
            "Página",
            "Situação da amostra",
            "Links na amostra",
            "Links encontrados",
            "Fração da amostra",
            "Mecanismos com correspondência",
            "Página lida em (UTC)",
        ),
        [],
    )
    gaps: Sheet = (("Chave", "Página", "Situação da amostra", "Link sem correspondência"), [])
    complements: Sheet = (
        (
            "Chave",
            "Página",
            "Situação da amostra",
            "Ordem sugerida",
            "Mecanismo",
            "Links correspondentes",
            "Links acrescentados",
            "Leitura parcial",
            "Mecanismo lido em (UTC)",
            "Página lida em (UTC)",
        ),
        [],
    )
    audits = {key: {s.url: s for s in values} for key, values in (checks or {}).items()}
    review: Sheet = (
        (
            "Chave",
            "Fonte",
            "Entrada",
            "Estado",
            "Grupo de revisão",
            "Motivo",
            "Páginas lidas na fonte",
            "Falha transitória",
            "Próximo passo sugerido",
        ),
        [],
    )

    def mechanism_row(key: str, page: str, url: str, label: str, previous: str) -> tuple[Cell, ...]:
        audit = audits.get(key, {}).get(url)
        return (
            key,
            page,
            url,
            label,
            audit.verdict + "; conferência da amostra" if audit else previous,
            audit.checked_at.isoformat() if audit else None,
            len(audit.listed_urls) if audit else None,
            "sim" if audit and audit.partial else "não" if audit else None,
            audit.detail if audit else None,
        )

    shown: set[tuple[str, str]] = set()
    for site in sites:
        known = pages.get(site.key, [])
        index = CoverageIndex(list(audits.get(site.key, {}).values()))
        measured_pages = {p.url: index.compare(p) for p in known if p.is_section_candidate}
        measured_counts = Counter(m.status for m in measured_pages.values())
        detail = next((p.detail for p in known if p.kind == "root"), "")
        progress = summarize_map(site.key, known, detail)
        state = (
            "não iniciada"
            if not known
            else "pendente com falhas"
            if progress.pending and progress.issues
            else "pendente"
            if progress.pending
            else "com falhas"
            if progress.issues or detail
            else "fila explorada"
        )
        primary = previous.get(site.key)
        summary[1].append(
            (
                site.name,
                site.key,
                site.country,
                state,
                progress.sections,
                progress.visited,
                progress.pending,
                progress.issues,
                progress.excluded,
                len({u for p in known for u in p.feeds + p.feed_indexes}),
                "não validada",
                primary.url if primary else None,
                detail,
                sum(m.sampled > 0 for m in measured_pages.values()),
                measured_counts["sample_full"],
                measured_counts["partial"],
                measured_counts["no_match"],
                measured_counts["no_sample"],
                len(audits.get(site.key, {})),
                sum(p.page_type == "article" for p in known),
                sum(p.page_type == "service" for p in known),
                sum(p.page_type == "section" for p in known),
                sum(p.is_section_candidate and p.discovery_basis == "navigation" for p in known),
                sum(
                    p.is_section_candidate and p.discovery_basis == "mechanism_path" for p in known
                ),
            )
        )
        for p in known:
            checked = p.checked_at.isoformat() if p.checked_at else None
            issue = review_entry(p)
            if issue is not None:
                review[1].append(
                    (
                        site.key,
                        site.name,
                        p.url,
                        p.status.value,
                        issue.category,
                        p.detail,
                        progress.visited,
                        "sim" if issue.transient else "não",
                        issue.action,
                    )
                )
            inventory[1].append(
                (
                    site.key,
                    site.name,
                    p.url,
                    p.found_on,
                    p.kind,
                    p.status.value,
                    p.title,
                    len(p.article_urls),
                    p.attempts,
                    checked,
                    p.detail,
                    "não validada",
                    p.page_type if p.kind != "taxonomy" else "taxonomy",
                    p.type_basis,
                    p.discovery_basis,
                )
            )
            for label, urls in (("feed/sitemap", p.feeds), ("índice de feeds", p.feed_indexes)):
                sources[1].extend(
                    mechanism_row(site.key, p.url, url, label, "pendente") for url in urls
                )
                shown.update((site.key, url) for url in urls)
            samples[1].extend((site.key, p.url, url) for url in p.article_urls)
            if p.is_section_candidate:
                measured = measured_pages[p.url]
                coverage[1].append(
                    (
                        site.key,
                        p.url,
                        measured.status,
                        measured.sampled,
                        measured.matched,
                        measured.matched / measured.sampled if measured.sampled else None,
                        "\n".join(measured.sources),
                        checked,
                    )
                )
                gaps[1].extend((site.key, p.url, measured.status, url) for url in measured.missing)
                complements[1].extend(
                    (
                        site.key,
                        p.url,
                        measured.status,
                        order,
                        contribution.source.url,
                        contribution.matched,
                        contribution.added,
                        "sim" if contribution.source.partial else "não",
                        contribution.source.checked_at.isoformat(),
                        checked,
                    )
                    for order, contribution in enumerate(measured.suggested, 1)
                )
    selected = {site.key for site in sites}
    for key, candidate, _ in candidates:
        if key in selected:
            kind = candidate.kind.value if candidate.kind else "desconhecido"
            sources[1].append(
                mechanism_row(
                    key,
                    "",
                    candidate.url,
                    f"sondagem anterior: {kind}",
                    f"{candidate.verdict.value}; cobertura por seção pendente",
                )
            )
            shown.add((key, candidate.url))
    for key, values in audits.items():
        if key in selected:
            for source in values.values():
                if (key, source.url) not in shown:
                    sources[1].append(mechanism_row(key, "", source.url, "auditoria do mapa", ""))
    return {
        "Fontes": summary,
        "Páginas": inventory,
        "Mecanismos": sources,
        "Amostras": samples,
        "Cobertura": coverage,
        "Lacunas": gaps,
        "Complementos": complements,
        "Revisão": review,
    }


def write_mapping(tables: Mapping[str, Sheet], path: Path) -> None:
    """Exporta XLSX ou um CSV por aba, protegendo texto de terceiros contra fórmulas."""
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == ".csv":
        for i, (name, (headers, rows)) in enumerate(tables.items()):
            target = path if i == 0 else path.with_name(f"{path.stem}-{name}{path.suffix}")
            with target.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(headers)
                writer.writerows(tuple(_cell(v) for v in row) for row in rows)
        return
    if path.suffix.lower() != ".xlsx":
        raise ValueError("a saída precisa terminar em .xlsx ou .csv")
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as error:
        raise RuntimeError("para gerar .xlsx instale o extra: pip install '.[xlsx]'") from error
    book = Workbook(write_only=True)
    for name, (headers, rows) in _sheet_parts(tables):
        sheet = book.create_sheet(name)
        sheet.freeze_panes = "C2"
        for index, header in enumerate(headers, 1):
            width = min(30, max(16, len(header) + 2))
            if header in (
                "Página",
                "Entrada",
                "Encontrada em",
                "Mecanismo",
                "Principal anterior",
            ) or (header.startswith("Link ")):
                width = 65
            elif header in (
                "Chave",
                "Fonte",
                "Título",
                "Detalhe",
                "Observação",
                "Motivo",
                "Próximo passo sugerido",
            ):
                width = 40
            sheet.column_dimensions[_column(index)].width = width
        sheet.row_dimensions[1].height = 32
        cells = [WriteOnlyCell(sheet, value=h) for h in headers]
        for cell in cells:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.append(cells)
        for row in rows:
            sheet.append([_cell(v) for v in row])
        sheet.auto_filter.ref = f"A1:{_column(len(headers))}{len(rows) + 1}"
    with NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        book.save(temporary)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def _column(n: int) -> str:
    """Nome Excel de uma coluna (base 26)."""
    result = ""
    while n:
        n, r = divmod(n - 1, 26)
        result = chr(65 + r) + result
    return result


def _cell(value: Cell | datetime) -> Cell:
    """Converte célula preservando texto, inclusive quando começa por sinal de fórmula."""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", "", value)
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value
